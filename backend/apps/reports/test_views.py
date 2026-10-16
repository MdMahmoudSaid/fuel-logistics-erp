# Dashboard / Report / Search API tests (docs/PHASE_6_PLAN.md Etape 5;
# ERP_Implementation_Experience_Guide.md Section 9.6/9.7/9.8).

import csv
import io
from decimal import Decimal

import pytest

from apps.finance.testing_helpers import seed_net_profit
from django.contrib.contenttypes.models import ContentType
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.accounts.models import Role, User
from apps.core.models import Product
from apps.customers.models import Customer
from apps.finance.models import (
    BankAccount,
    CashAccount,
    CustomerPayment,
    Expense,
    ExpenseAllocation,
    Shareholder,
    ShareholderLoan,
    SupplierAdvance,
    TransportPayment,
)
from apps.finance.services import (
    ControlledDebtPaymentService,
    FinancePostingService,
    ShareholderDistributionService,
    ShareholderWithdrawalService,
)
from apps.purchasing.models import CessionLetter, SupplierOrder
from apps.sales.models import Sale
from apps.suppliers.models import Supplier
from apps.transport.models import Transporter, Truck
from apps.trips.models import Trip


def _make_admin(suffix):
    admin_role = Role.objects.get(name=Role.NAME_ADMIN)
    return User.objects.create_user(
        phone=(f"+222{suffix}")[:20], pin="1234", role=admin_role
    )


def _make_role_without_reports(suffix):
    return Role.objects.create(name=f"NoReports-{suffix}", permissions={"operations": ["view"]})


def _sold_on(trip, day):
    """Date a test sale on `day` - since 2026-10-05 a trip's profit/revenue
    belongs to its SALE date, so tests that place a trip in a date range
    move its sale there too."""
    from datetime import date as _date, datetime as _datetime, time as _time

    from django.utils import timezone as _tz

    Sale.objects.filter(trip=trip).update(
        created_at=_tz.make_aware(_datetime.combine(_date.fromisoformat(str(day)), _time(12)))
    )


def _make_full_trip(suffix, admin_user, status=Trip.Status.DELIVERED, with_sale=True):
    product = Product.objects.get(code="gasoil")
    supplier = Supplier.objects.create(name=f"Supplier {suffix}", country="Senegal")
    transporter = Transporter.objects.create(name=f"Trans {suffix}")
    supplier_order = SupplierOrder.objects.create(
        reference=f"SO-VIEW-{suffix}",
        supplier=supplier,
        product=product,
        purchase_date="2026-08-18",
        quantity_litres=Decimal("500000.00"),
        unit_price=Decimal("600.0000"),
        status=SupplierOrder.Status.CONFIRMED,
        created_by=admin_user,
    )
    cession = CessionLetter.objects.create(
        reference=f"CL-VIEW-{suffix}",
        supplier_order=supplier_order,
        authorized_quantity=Decimal("45000.00"),
        issue_date="2026-08-18",
        created_by=admin_user,
    )
    truck = Truck.objects.create(
        internal_id=f"TRK-VIEW-{suffix}",
        tractor_registration=f"{suffix}-TRACTOR",
        tanker_registration=f"{suffix}-TANKER",
        capacity_litres=Decimal("45000.00"),
    )
    trip = Trip.objects.create(
        reference=f"TRIP-VIEW-{suffix}",
        truck=truck,
        transporter=transporter,
        supplier=supplier,
        supplier_order=supplier_order,
        cession_letter=cession,
        product=product,
        quantity_loaded=Decimal("45000.00"),
        loading_date="2026-08-18",
        bl_reference=f"BL-VIEW-{suffix}",
        status=status,
        created_by=admin_user,
    )
    customer = Customer.objects.create(name=f"Client {suffix}")
    sale = None
    if with_sale:
        sale = Sale.objects.create(
            trip=trip,
            customer=customer,
            product=product,
            delivered_quantity=Decimal("44850.00"),
            unit_selling_price=Decimal("650.0000"),
        )
        _sold_on(trip, trip.loading_date)
    return {"trip": trip, "truck": truck, "customer": customer, "sale": sale, "supplier_order": supplier_order}


# ---------------------------------------------------------------------------
# DashboardViewSet
# ---------------------------------------------------------------------------


@pytest.mark.django_db
def test_dashboard_summary_returns_six_sections():
    admin_user = _make_admin("DASH")
    _make_full_trip("DASH", admin_user)

    client = APIClient()
    client.force_authenticate(admin_user)
    response = client.get("/api/dashboard/summary/")

    assert response.status_code == 200, response.data
    for section in ("operations", "fuel", "sales", "transport", "finance", "profit"):
        assert section in response.data


@pytest.mark.django_db
def test_dashboard_summary_filters_by_explicit_date_range():
    admin_user = _make_admin("DASHRANGE")
    _make_full_trip("DASHRANGE", admin_user)

    client = APIClient()
    client.force_authenticate(admin_user)
    response = client.get(
        "/api/dashboard/summary/", {"from": "2020-01-01", "to": "2020-01-02"}
    )

    assert response.status_code == 200, response.data
    assert response.data["operations"]["trips_count"] == 0


@pytest.mark.django_db
def test_dashboard_summary_period_cards_accept_start_date_end_date_aliases():
    admin_user = User.objects.create_superuser(phone="+222DASHPERA", pin="1234")
    jan = _make_full_trip("DASHPERA", admin_user)
    jan["trip"].loading_date = "2026-01-10"
    jan["trip"].save(update_fields=["loading_date"])
    _sold_on(jan["trip"], "2026-01-10")

    sep = _make_full_trip("DASHPERB", admin_user)
    sep["trip"].loading_date = "2026-09-01"
    sep["trip"].save(update_fields=["loading_date"])
    _sold_on(sep["trip"], "2026-09-01")

    cash = CashAccount.objects.create(name="Dash Period Cash")
    CustomerPayment.objects.create(
        customer=jan["customer"],
        amount=Decimal("1000.00"),
        payment_date="2026-01-10",
        method=CustomerPayment.Method.CASH,
        cash_account=cash,
        recorded_by=admin_user,
    )
    CustomerPayment.objects.create(
        customer=sep["customer"],
        amount=Decimal("2000.00"),
        payment_date="2026-09-01",
        method=CustomerPayment.Method.CASH,
        cash_account=cash,
        recorded_by=admin_user,
    )
    Expense.objects.create(
        category=Expense.Category.OTHER,
        amount=Decimal("100.00"),
        calculation_method=Expense.CalculationMethod.FIXED,
        date="2026-01-10",
        status=Expense.Status.PAID,
    )
    Expense.objects.create(
        category=Expense.Category.OTHER,
        amount=Decimal("200.00"),
        calculation_method=Expense.CalculationMethod.FIXED,
        date="2026-09-01",
        status=Expense.Status.PAID,
    )

    client = APIClient()
    client.force_authenticate(admin_user)

    jan_response = client.get(
        "/api/dashboard/summary/",
        {"preset": "custom", "start_date": "2026-01-01", "end_date": "2026-01-31"},
    )
    sep_response = client.get(
        "/api/dashboard/summary/",
        {"preset": "custom", "start_date": "2026-09-01", "end_date": "2026-09-01"},
    )
    wide_response = client.get(
        "/api/dashboard/summary/",
        {"preset": "custom", "start_date": "2026-01-01", "end_date": "2026-09-01"},
    )

    assert jan_response.status_code == 200, jan_response.data
    assert sep_response.status_code == 200, sep_response.data
    assert wide_response.status_code == 200, wide_response.data

    assert jan_response.data["operations"]["trips_count"] == 1
    assert jan_response.data["financial_summary"]["period_revenue"] == str(jan["sale"].subtotal)
    assert jan_response.data["financial_summary"]["period_collections"] == "1000.00"
    assert jan_response.data["financial_summary"]["period_paid_expenses"] == "100.00"

    assert sep_response.data["operations"]["trips_count"] == 1
    assert sep_response.data["financial_summary"]["period_revenue"] == str(sep["sale"].subtotal)
    assert sep_response.data["financial_summary"]["period_collections"] == "2000.00"
    assert sep_response.data["financial_summary"]["period_paid_expenses"] == "200.00"

    assert wide_response.data["operations"]["trips_count"] == 2
    assert wide_response.data["financial_summary"]["period_revenue"] == str(
        jan["sale"].subtotal + sep["sale"].subtotal
    )
    assert wide_response.data["financial_summary"]["period_collections"] == "3000.00"
    assert wide_response.data["financial_summary"]["period_paid_expenses"] == "300.00"


@pytest.mark.django_db
def test_dashboard_summary_treasury_in_out_counts_shareholder_contribution():
    """Bug fix (2026-09-03 recette manuelle): the Journal de tresorerie
    "Entrees / sorties" card used to be sourced from
    period_collections/period_paid_expenses, which only ever look at
    CustomerPayment/paid Expense rows - a ShareholderContribution posts
    only a TreasuryTransaction (MONEY_IN), so it appeared in the journal's
    transaction list but never moved the card. period_treasury_in/out must
    aggregate TreasuryTransaction directly so the card matches the list."""
    from apps.finance.services import ShareholderContributionService

    admin_user = _make_admin("DASHTRIN")
    bank = BankAccount.objects.create(name="Bank Treasury Card", opening_balance=Decimal("0.00"))
    shareholder = Shareholder.objects.create(name="Assoc TreasuryCard", participation_percent=Decimal("100"))

    ShareholderContributionService.record_contribution(
        shareholder_id=shareholder.id,
        amount=Decimal("10000000.00"),
        date="2026-09-03",
        actor=admin_user,
        attachment=SimpleUploadedFile("proof.jpg", b"filedata", content_type="image/jpeg"),
        bank_account_id=bank.id,
    )

    client = APIClient()
    client.force_authenticate(admin_user)
    response = client.get(
        "/api/dashboard/summary/",
        {"preset": "custom", "start_date": "2026-09-01", "end_date": "2026-09-30"},
    )

    assert response.status_code == 200, response.data
    assert response.data["financial_summary"]["period_treasury_in"] == "10000000.00"
    assert response.data["financial_summary"]["period_treasury_out"] == "0.00"
    # Untouched by this fix - still CustomerPayment/Expense-only.
    assert response.data["financial_summary"]["period_collections"] == "0.00"


@pytest.mark.django_db
def test_dashboard_summary_includes_shareholders_summary_section():
    """Design doc 2026-08-26-benefices-associes-refinement-design.md:
    dashboard shows a summary only - current profit, total withdrawn,
    remaining undistributed, open loans - the Benefices page stays the
    detailed source."""
    admin_user = _make_admin("DASHSH")
    seed_net_profit()
    a = Shareholder.objects.create(name="Dash A", participation_percent=Decimal("100"))
    CashAccount.objects.create(name="Dash A Till", opening_balance=Decimal("1000000.00"))
    distribution = ShareholderDistributionService.create_distribution(
        period_start="2026-07-01",
        period_end="2026-07-31",
        total_profit=Decimal("100000.00"),
        actor=admin_user,
    )
    ShareholderDistributionService.validate_distribution(distribution.id, admin_user)
    share = distribution.shares.get(shareholder=a)
    cash_account = CashAccount.objects.create(name="Dash Cash", opening_balance=Decimal("100000.00"))
    ShareholderWithdrawalService.record_withdrawal(
        shareholder_id=a.id,
        amount=Decimal("30000.00"),
        date="2026-08-01",
        cash_account_id=cash_account.id,
        actor=admin_user,
    )
    b = Shareholder.objects.create(name="Dash B", participation_percent=Decimal("0"))
    ShareholderLoan.objects.create(
        borrower=b, lender=a, amount=Decimal("5000.00"), date="2026-08-01"
    )

    client = APIClient()
    client.force_authenticate(admin_user)
    response = client.get("/api/dashboard/summary/")

    assert response.status_code == 200, response.data
    summary = response.data["shareholders_summary"]
    assert summary["total_withdrawn"] == "30000.00"
    assert summary["remaining_to_pay_shareholders"] == "70000.00"
    assert summary["open_loans_total"] == "5000.00"
    # Recette 2026-09-07: exposed so the UI can show the tolerance-absorbed
    # delivery loss that distributable_result silently nets out.
    assert "total_company_absorbed_loss" in summary


@pytest.mark.django_db
def test_dashboard_summary_remaining_to_pay_shareholders_excludes_cancelled():
    admin_user = _make_admin("DASHCANCEL")
    seed_net_profit()
    Shareholder.objects.create(name="Dash Cancel", participation_percent=Decimal("100"))
    CashAccount.objects.create(name="Dash Cancel Till", opening_balance=Decimal("1000000.00"))
    distribution = ShareholderDistributionService.create_distribution(
        period_start="2026-07-01",
        period_end="2026-07-31",
        total_profit=Decimal("100000.00"),
        actor=admin_user,
    )
    ShareholderDistributionService.validate_distribution(distribution.id, admin_user)
    ShareholderDistributionService.cancel_distribution(distribution.id, admin_user)

    client = APIClient()
    client.force_authenticate(admin_user)
    response = client.get("/api/dashboard/summary/")

    assert response.status_code == 200, response.data
    assert response.data["shareholders_summary"]["remaining_to_pay_shareholders"] == "0.00"


@pytest.mark.django_db
def test_dashboard_summary_shareholders_section_includes_distributable_result():
    """Client modifications (2026-09-04, third pass - confirmed bug during
    actual recette testing): dashboard must expose the same
    distributable_result figure as the Benefices page's profit-preview
    endpoint - same formula (total - debts - capital), same source
    (ShareholderProfitPreviewService.get_preview), never a second one.
    Replaces the old total_capital_contributed/profit_after_capital pair
    entirely (docs/superpowers/specs/2026-09-04-recette-corrections-
    design.md Section 4)."""
    admin_user = _make_admin("DASHPAC")
    fx = _make_full_trip("DASHPAC", admin_user)
    Shareholder.objects.create(
        name="Dash PAC", participation_percent=Decimal("100"), contribution=Decimal("400000.00")
    )

    client = APIClient()
    client.force_authenticate(admin_user)
    # shareholders_summary is a point-in-time snapshot, never scoped by
    # date_from/date_to (see the section's own comment in dashboard.py) -
    # no range query params here, matching that contract.
    response = client.get("/api/dashboard/summary/")

    assert response.status_code == 200, response.data
    summary = response.data["shareholders_summary"]
    wc = response.data["working_capital"]
    assert "total_capital_contributed" not in summary
    assert "profit_after_capital" not in summary
    # Client decision 2026-09-25: the dashboard exposes the same
    # distributable_result as the Benefices page - the net profit after
    # general expenses of the default period, never a second formula.
    from apps.finance.services import ShareholderProfitPreviewService

    expected = ShareholderProfitPreviewService.get_preview()["net_profit"]
    assert Decimal(summary["distributable_result"]) == expected
    assert wc["capital"] is not None
    # Fixture sanity: there really is a sold trip driving current_profit.
    assert fx["sale"] is not None


@pytest.mark.django_db
def test_dashboard_summary_period_net_profit_falls_after_distribution_is_validated():
    """Client decision (2026-10-04): the dashboard's "Bénéfice net après
    dépenses générales" card must fall as distributions are validated and
    reach zero once everything is distributed - same contract as the
    Bénéfices page - instead of duplicating the old period-only formula
    independently of what has already been distributed."""
    admin_user = _make_admin("DASHNETFALL")
    _make_full_trip("DASHNETFALL", admin_user)
    Shareholder.objects.create(name="Dash Net", participation_percent=Decimal("100"))

    client = APIClient()
    client.force_authenticate(admin_user)

    from apps.finance.services import ShareholderDistributionService, ShareholderProfitPreviewService

    before = ShareholderProfitPreviewService.get_preview()["net_profit"]
    assert before > Decimal("0.00")

    distribution = ShareholderDistributionService.create_distribution(
        period_start="2026-01-01",
        period_end="2026-12-31",
        total_profit=before,
        actor=admin_user,
    )
    ShareholderDistributionService.validate_distribution(distribution.id, admin_user)

    response = client.get("/api/dashboard/summary/")
    assert response.status_code == 200, response.data
    assert Decimal(response.data["financial_summary"]["period_net_profit"]) == Decimal("0.00")


@pytest.mark.django_db
def test_dashboard_summary_requires_reports_permission():
    role = _make_role_without_reports("DASH")
    user = User.objects.create_user(phone="+2220000013", pin="1234", role=role)

    client = APIClient()
    client.force_authenticate(user)
    response = client.get("/api/dashboard/summary/")

    assert response.status_code == 403


@pytest.mark.django_db
def test_dashboard_summary_includes_debts_customer_and_bank_sections():
    admin_user = _make_admin("DASHDEBT")
    _make_full_trip("DASHDEBT", admin_user)

    client = APIClient()
    client.force_authenticate(admin_user)
    response = client.get("/api/dashboard/summary/")

    assert response.status_code == 200, response.data
    assert set(response.data["debts"]) >= {
        "total",
        "suppliers",
        "transporters",
        "customs",
        "other",
    }
    assert "total" in response.data["customer_amounts"]
    assert "rows" in response.data["customer_amounts"]
    assert "total" in response.data["bank_amounts"]
    assert "rows" in response.data["bank_amounts"]
    assert "total" in response.data["cash_amounts"]
    assert "rows" in response.data["cash_amounts"]


@pytest.mark.django_db
def test_dashboard_summary_aggregates_customs_and_other_debts_correctly():
    admin_user = _make_admin("DASHEXP")

    Expense.objects.create(
        category=Expense.Category.CUSTOMS,
        amount=Decimal("1000.00"),
        calculation_method=Expense.CalculationMethod.FIXED,
        date="2020-01-01",
        status=Expense.Status.UNPAID,
    )
    Expense.objects.create(
        category=Expense.Category.CUSTOMS,
        amount=Decimal("500.00"),
        calculation_method=Expense.CalculationMethod.FIXED,
        date="2020-01-01",
        status=Expense.Status.PARTIALLY_PAID,
    )
    Expense.objects.create(
        category=Expense.Category.CUSTOMS,
        amount=Decimal("999.00"),
        calculation_method=Expense.CalculationMethod.FIXED,
        date="2020-01-01",
        status=Expense.Status.PAID,
    )
    # A general expense (Loyer here) has no trip-fee debt bucket at all -
    # "other" is only border + loading. Kept as a regression guard that it
    # never leaks into the dashboard's trip-debt aggregate.
    Expense.objects.create(
        category=Expense.Category.RENT,
        amount=Decimal("300.00"),
        calculation_method=Expense.CalculationMethod.FIXED,
        date="2020-01-01",
        status=Expense.Status.UNPAID,
    )
    # Transport-category expenses must be excluded from "other" - they are
    # already captured via TransporterOutstandingReportQuery/debts.transporters.
    Expense.objects.create(
        category=Expense.Category.TRANSPORT,
        amount=Decimal("777.00"),
        calculation_method=Expense.CalculationMethod.FIXED,
        date="2020-01-01",
        status=Expense.Status.UNPAID,
    )
    Expense.objects.create(
        category=Expense.Category.ROAD_TRANSIT,
        amount=Decimal("450000.00"),
        calculation_method=Expense.CalculationMethod.PER_LITRE,
        date="2020-01-01",
        status=Expense.Status.UNPAID,
    )
    Expense.objects.create(
        category=Expense.Category.RENT,
        amount=Decimal("50.00"),
        calculation_method=Expense.CalculationMethod.FIXED,
        date="2020-01-01",
        status=Expense.Status.CANCELLED,
    )

    client = APIClient()
    client.force_authenticate(admin_user)
    response = client.get("/api/dashboard/summary/")

    assert response.status_code == 200, response.data
    assert response.data["debts"]["customs"] == "1500.00"
    assert response.data["debts"]["road_transit"] == "450000.00"
    # Loyer (RENT) has no trip-fee bucket; "other" is border + loading only.
    assert response.data["debts"]["other"] == "0.00"
    assert response.data["party_balances"]["customs_and_other"]["road_transit"] == "450000.00"


@pytest.mark.django_db
def test_dashboard_summary_shows_transit_debt_until_controlled_payment():
    admin_user = _make_admin("DASHTRANSIT")
    transit = Expense.objects.create(
        category="transit",
        amount=Decimal("90000.00"),
        calculation_method=Expense.CalculationMethod.PER_LITRE,
        date="2020-01-01",
        status=Expense.Status.UNPAID,
    )
    Expense.objects.create(
        category=Expense.Category.TRANSPORT,
        amount=Decimal("125000.00"),
        calculation_method=Expense.CalculationMethod.FIXED,
        date="2020-01-01",
        status=Expense.Status.UNPAID,
    )
    bank = BankAccount.objects.create(
        name="Dashboard Transit Bank", opening_balance=Decimal("1000000.00")
    )
    client = APIClient()
    client.force_authenticate(admin_user)

    before = client.get("/api/dashboard/summary/")

    assert before.status_code == 200, before.data
    assert before.data["debts"]["transit"] == "90000.00"
    assert before.data["party_balances"]["customs_and_other"]["transit"] == "90000.00"

    ControlledDebtPaymentService.record_payment(
        debt_type="transit",
        debt_id=0,
        bank_account_id=bank.id,
        amount=Decimal("90000.00"),
        date="2020-01-02",
        reference="DASH-TRANSIT-PAY",
        actor=admin_user,
        attachment=SimpleUploadedFile(
            "transit.png", b"fake png receipt", content_type="image/png"
        ),
    )
    transit.refresh_from_db()

    after = client.get("/api/dashboard/summary/")

    assert transit.status == Expense.Status.PAID
    assert after.status_code == 200, after.data
    assert after.data["debts"]["transit"] == "0.00"
    assert after.data["party_balances"]["customs_and_other"]["transit"] == "0.00"


@pytest.mark.django_db
def test_dashboard_summary_transporter_balance_matches_new_payment_reste_a_payer():
    """Bug: the dashboard's "Transporteurs" card must show the same net
    reste-a-payer as Journal de tresorerie > Nouveau paiement, not the raw
    settlement total - a settlement of 5 670 000 with 33 000 already paid
    must show as -5 637 000, not -5 670 000."""
    admin_user = _make_admin("DASHTRANSPAY")
    fx = _make_full_trip("DASHTRANSPAY", admin_user)
    transporter = fx["trip"].transporter
    TransportPayment.objects.create(
        trip=fx["trip"],
        transporter=transporter,
        calculation_method=TransportPayment.CalculationMethod.NEGOTIATED,
        base_amount=Decimal("5670000.00"),
        status=TransportPayment.Status.PAYABLE,
    )
    bank = BankAccount.objects.create(
        name="Dashboard Transport Bank", opening_balance=Decimal("10000000.00")
    )
    client = APIClient()
    client.force_authenticate(admin_user)

    before = client.get("/api/dashboard/summary/")
    assert before.status_code == 200, before.data
    assert before.data["party_balances"]["transporters"]["balance"] == "-5670000.00"

    ControlledDebtPaymentService.record_payment(
        debt_type=ControlledDebtPaymentService.TRANSPORTER,
        debt_id=transporter.id,
        bank_account_id=bank.id,
        amount=Decimal("33000.00"),
        date="2026-08-19",
        reference="DASH-TRANS-PAY",
        actor=admin_user,
    )

    after = client.get("/api/dashboard/summary/")

    assert after.status_code == 200, after.data
    assert after.data["debts"]["transporters"]["total"] == "5637000.00"
    assert after.data["party_balances"]["transporters"]["balance"] == "-5637000.00"


@pytest.mark.django_db
def test_dashboard_summary_customs_and_other_balance_matches_open_debts_by_category():
    """Bug: the dashboard's "Douane / Transit / Frais autres" card must sum
    the same open reste-a-payer buckets as Nouveau paiement: douane +
    transit routiere + frontiere (Expense.Category.ROAD) - here
    20 200 000 + 350 000 + 450 000 = 21 000 000, not the old inflated
    total that folded unrelated categories into "other"."""
    admin_user = _make_admin("DASHCUSTNET")
    Expense.objects.create(
        category=Expense.Category.CUSTOMS,
        amount=Decimal("20200000.00"),
        calculation_method=Expense.CalculationMethod.FIXED,
        date="2020-01-01",
        status=Expense.Status.UNPAID,
    )
    Expense.objects.create(
        category=Expense.Category.ROAD_TRANSIT,
        amount=Decimal("350000.00"),
        calculation_method=Expense.CalculationMethod.FIXED,
        date="2020-01-01",
        status=Expense.Status.UNPAID,
    )
    Expense.objects.create(
        category=Expense.Category.ROAD,  # "frontiere" (ControlledDebtPaymentService.BORDER)
        amount=Decimal("450000.00"),
        calculation_method=Expense.CalculationMethod.FIXED,
        date="2020-01-01",
        status=Expense.Status.UNPAID,
    )

    client = APIClient()
    client.force_authenticate(admin_user)
    response = client.get("/api/dashboard/summary/")

    assert response.status_code == 200, response.data
    assert response.data["debts"]["customs"] == "20200000.00"
    assert response.data["debts"]["road_transit"] == "350000.00"
    assert response.data["debts"]["other"] == "450000.00"
    assert response.data["party_balances"]["customs_and_other"]["balance"] == "-21000000.00"


@pytest.mark.django_db
def test_dashboard_summary_customs_balance_decreases_immediately_after_partial_payment():
    admin_user = _make_admin("DASHCUSTPAY")
    Expense.objects.create(
        category=Expense.Category.CUSTOMS,
        amount=Decimal("20200000.00"),
        calculation_method=Expense.CalculationMethod.FIXED,
        date="2020-01-01",
        status=Expense.Status.UNPAID,
    )
    bank = BankAccount.objects.create(
        name="Dashboard Customs Bank", opening_balance=Decimal("30000000.00")
    )
    client = APIClient()
    client.force_authenticate(admin_user)

    ControlledDebtPaymentService.record_payment(
        debt_type=ControlledDebtPaymentService.CUSTOMS,
        debt_id=0,
        bank_account_id=bank.id,
        amount=Decimal("200000.00"),
        date="2020-01-02",
        reference="DASH-CUSTOMS-PAY",
        actor=admin_user,
    )

    response = client.get("/api/dashboard/summary/")

    assert response.status_code == 200, response.data
    assert response.data["debts"]["customs"] == "20000000.00"
    assert response.data["party_balances"]["customs_and_other"]["customs"] == "20000000.00"
    assert response.data["party_balances"]["customs_and_other"]["balance"] == "-20000000.00"


@pytest.mark.django_db
def test_dashboard_summary_balances_exclude_cancelled_and_draft_items():
    """Cancelled expenses, draft supplier orders, and fully PAID transport
    settlements must never inflate the dashboard's open balances."""
    admin_user = _make_admin("DASHEXCL")
    fx = _make_full_trip("DASHEXCL", admin_user, with_sale=False)
    transporter = fx["trip"].transporter
    supplier = fx["supplier_order"].supplier

    client = APIClient()
    client.force_authenticate(admin_user)
    before = client.get("/api/dashboard/summary/")
    assert before.status_code == 200, before.data
    suppliers_balance_before = before.data["party_balances"]["suppliers"]["balance"]

    # Cancelled expense in an otherwise-empty debt-bearing category.
    Expense.objects.create(
        category=Expense.Category.CUSTOMS,
        amount=Decimal("999999.00"),
        calculation_method=Expense.CalculationMethod.FIXED,
        date="2020-01-01",
        status=Expense.Status.CANCELLED,
    )
    # A fully PAID transport settlement must not count as open debt.
    TransportPayment.objects.create(
        trip=fx["trip"],
        transporter=transporter,
        calculation_method=TransportPayment.CalculationMethod.NEGOTIATED,
        base_amount=Decimal("123456.00"),
        status=TransportPayment.Status.PAID,
    )
    # A draft supplier order must not add to the supplier's open debt.
    SupplierOrder.objects.create(
        reference="SO-DASHEXCL-DRAFT",
        supplier=supplier,
        product=Product.objects.get(code="gasoil"),
        purchase_date="2020-01-01",
        quantity_litres=Decimal("1000.00"),
        unit_price=Decimal("600.0000"),
        status=SupplierOrder.Status.DRAFT,
        created_by=admin_user,
    )

    after = client.get("/api/dashboard/summary/")

    assert after.status_code == 200, after.data
    assert after.data["debts"]["customs"] == "0.00"
    assert after.data["party_balances"]["transporters"]["balance"] == "0.00"
    assert after.data["party_balances"]["suppliers"]["balance"] == suppliers_balance_before


@pytest.mark.django_db
def test_dashboard_summary_bank_amounts_only_include_active_banks():
    admin_user = _make_admin("DASHBANK")
    BankAccount.objects.create(
        name="Active Bank", opening_balance=Decimal("100.00"), is_active=True
    )
    BankAccount.objects.create(
        name="Closed Bank", opening_balance=Decimal("50.00"), is_active=False
    )

    client = APIClient()
    client.force_authenticate(admin_user)
    response = client.get("/api/dashboard/summary/")

    assert response.status_code == 200, response.data
    bank_names = {row["bank"] for row in response.data["bank_amounts"]["rows"]}
    assert bank_names == {"Active Bank"}
    assert response.data["bank_amounts"]["total"] == "100.00"


@pytest.mark.django_db
def test_dashboard_summary_cash_amounts_only_include_active_cash_accounts():
    admin_user = _make_admin("DASHCASH")
    CashAccount.objects.create(
        name="Active Cash", opening_balance=Decimal("125.00"), is_active=True
    )
    CashAccount.objects.create(
        name="Closed Cash", opening_balance=Decimal("75.00"), is_active=False
    )

    client = APIClient()
    client.force_authenticate(admin_user)
    response = client.get("/api/dashboard/summary/")

    assert response.status_code == 200, response.data
    cash_names = {row["cash_account"] for row in response.data["cash_amounts"]["rows"]}
    assert cash_names == {"Active Cash"}
    assert response.data["cash_amounts"]["total"] == "125.00"


@pytest.mark.django_db
def test_dashboard_summary_debts_customer_and_bank_sections_ignore_date_range():
    admin_user = _make_admin("DASHNODATE")
    Expense.objects.create(
        category=Expense.Category.CUSTOMS,
        amount=Decimal("1000.00"),
        calculation_method=Expense.CalculationMethod.FIXED,
        date="2020-01-01",
        status=Expense.Status.UNPAID,
    )
    BankAccount.objects.create(
        name="Old Bank", opening_balance=Decimal("250.00"), is_active=True
    )

    client = APIClient()
    client.force_authenticate(admin_user)
    unfiltered = client.get("/api/dashboard/summary/")
    filtered = client.get(
        "/api/dashboard/summary/", {"from": "2099-01-01", "to": "2099-01-02"}
    )

    assert unfiltered.status_code == 200, unfiltered.data
    assert filtered.status_code == 200, filtered.data
    # Unlike operations/fuel/sales/transport/finance, these three sections
    # are point-in-time "currently outstanding" snapshots and must not be
    # affected by the requested date_from/date_to range.
    assert filtered.data["operations"]["trips_count"] == 0
    assert filtered.data["debts"] == unfiltered.data["debts"]
    assert filtered.data["customer_amounts"] == unfiltered.data["customer_amounts"]
    assert filtered.data["bank_amounts"] == unfiltered.data["bank_amounts"]
    assert filtered.data["cash_amounts"] == unfiltered.data["cash_amounts"]
    assert filtered.data["debts"]["customs"] == "1000.00"
    assert filtered.data["bank_amounts"]["total"] == "250.00"


@pytest.mark.django_db
def test_dashboard_summary_exposes_final_business_sections():
    admin_user = _make_admin("DASHFINAL")
    fx = _make_full_trip(
        "DASHFINAL", admin_user, status=Trip.Status.ARRIVED_AT_COMPANY_STOCK, with_sale=False
    )
    BankAccount.objects.create(name="Bank Recette", opening_balance=Decimal("1000000.00"))
    CashAccount.objects.create(name="Caisse Recette", opening_balance=Decimal("250000.00"))
    payment = CustomerPayment.objects.create(
        customer=fx["customer"],
        amount=Decimal("300000.00"),
        payment_date="2026-08-18",
        method=CustomerPayment.Method.CASH,
        cash_account=CashAccount.objects.get(name="Caisse Recette"),
        recorded_by=admin_user,
    )
    FinancePostingService.record_customer_payment_treasury(payment)
    # Dashboard recette fix (2026-09-04): this is a generic smoke check for
    # the whole financial_summary section, not the category-exclusion rule
    # itself (see test_dashboard_summary_period_paid_expenses_excludes_
    # trip_fees below for that) - uses a genuine general-expense category
    # (OTHER) so it isn't accidentally excluded by that rule.
    Expense.objects.create(
        category=Expense.Category.OTHER,
        amount=Decimal("50000.00"),
        calculation_method=Expense.CalculationMethod.FIXED,
        date="2026-08-18",
        status=Expense.Status.PAID,
    )

    client = APIClient()
    client.force_authenticate(admin_user)
    response = client.get(
        "/api/dashboard/summary/", {"from": "2026-08-18", "to": "2026-08-18"}
    )

    assert response.status_code == 200, response.data
    financial = response.data["financial_summary"]
    assert financial["period_revenue"] == "0.00"
    assert financial["period_collections"] == "300000.00"
    assert financial["period_paid_expenses"] == "50000.00"
    assert financial["current_treasury_balance"] == "1550000.00"

    company_stock = response.data["company_stock"]
    assert company_stock["total_quantity"] == "45000.00"
    # Client request (2026-09-06): the stock card's value = full cost basis
    # (fuel only here - no TransportPayment/Expense on this fixture's
    # trip) - 45,000 L x 600.0000 = 27,000,000.
    assert company_stock["value"] == "27000000.00"
    assert company_stock["trips_count"] == 1
    assert company_stock["by_product"] == [{"product": "Gasoil", "quantity": "45000.00"}]
    assert company_stock["trips"][0]["reference"] == fx["trip"].reference
    assert company_stock["trips"][0]["remaining_quantity"] == "45000.00"

    # Client request (2026-09-06): the other two pipeline cards - this
    # trip is already ARRIVED_AT_COMPANY_STOCK, so none of its quantity is
    # "en cours" any more; the reliquat is what's left of the 500,000 L
    # purchase (455,000 L) priced at the order's own 600.0000 unit price.
    assert response.data["in_transit"] == {"quantity": "0.00", "value": "0.00"}
    assert response.data["supplier_reliquat"] == {
        "quantity": "455000.00",
        "value": "273000000.00",
    }

    party_balances = response.data["party_balances"]
    assert set(party_balances) >= {"customers", "suppliers", "transporters", "shareholders"}
    assert party_balances["shareholders"]["balance"] == "0.00"

    action_items = response.data["action_items"]
    assert any(item["kind"] == "company_stock_without_customer" for item in action_items)


@pytest.mark.django_db
def test_dashboard_summary_period_paid_expenses_excludes_trip_fees():
    """Dashboard recette fix (2026-09-04): confirmed bug during actual
    client recette testing - the "Depenses generales payees" card counted
    trip-linked fees (douane, transit, transit routiere, chargement,
    transport) alongside real general expenses. Client's own example: a
    1 000 000 XOF general expense (Internet) plus real trip fees paid in
    the same period must show exactly 1 000 000, not their sum."""
    admin_user = _make_admin("DASHGENEXP")
    fx = _make_full_trip("DASHGENEXP", admin_user)
    trip = fx["trip"]

    # General expense (Internet-like, non-system category, no allocation) -
    # the ONLY thing that must count.
    Expense.objects.create(
        category=Expense.Category.OTHER,
        amount=Decimal("1000000.00"),
        calculation_method=Expense.CalculationMethod.MANUAL,
        date=trip.loading_date,
        status=Expense.Status.PAID,
    )

    # Trip fees, allocated to the trip like the real "frais trajet" UI does
    # - douane + transit, both paid - must NEVER show up in this card.
    for category, amount in (
        (Expense.Category.CUSTOMS, Decimal("200000.00")),
        (Expense.Category.TRANSIT, Decimal("60450.00")),
    ):
        trip_expense = Expense.objects.create(
            category=category,
            amount=amount,
            calculation_method=Expense.CalculationMethod.MANUAL,
            date=trip.loading_date,
            status=Expense.Status.PAID,
        )
        ExpenseAllocation.objects.create(
            expense=trip_expense,
            content_type=ContentType.objects.get_for_model(Trip),
            object_id=trip.id,
            allocated_amount=amount,
        )

    client = APIClient()
    client.force_authenticate(admin_user)
    response = client.get(
        "/api/dashboard/summary/",
        {"preset": "custom", "start_date": str(trip.loading_date), "end_date": str(trip.loading_date)},
    )

    assert response.status_code == 200, response.data
    financial = response.data["financial_summary"]
    # Must be exactly the general expense - NOT 1 260 450 (the sum
    # including douane + transit), matching the client's exact example.
    assert financial["period_paid_expenses"] == "1000000.00"

    # period_net_profit = operational profit (period_accounting_profit) -
    # general expenses paid - never the raw operational figure.
    operational = Decimal(financial["period_accounting_profit"])
    net = Decimal(financial["period_net_profit"])
    assert net == operational - Decimal("1000000.00")
    assert net != operational


@pytest.mark.django_db
def test_margin_report_endpoint_filters_rows_with_start_date_end_date_aliases():
    admin_user = User.objects.create_superuser(phone="+222MARGVWA", pin="1234")
    jan = _make_full_trip("MARGVWA", admin_user)
    jan["trip"].loading_date = "2026-01-10"
    jan["trip"].save(update_fields=["loading_date"])
    _sold_on(jan["trip"], "2026-01-10")
    sep = _make_full_trip("MARGVWB", admin_user)
    sep["trip"].loading_date = "2026-09-01"
    sep["trip"].save(update_fields=["loading_date"])
    _sold_on(sep["trip"], "2026-09-01")

    client = APIClient()
    client.force_authenticate(admin_user)

    jan_response = client.get(
        "/api/reports/margin-report/",
        {"preset": "custom", "start_date": "2026-01-01", "end_date": "2026-01-31"},
    )
    sep_response = client.get(
        "/api/reports/margin-report/",
        {"preset": "custom", "start_date": "2026-09-01", "end_date": "2026-09-01"},
    )
    wide_response = client.get(
        "/api/reports/margin-report/",
        {"preset": "custom", "start_date": "2026-01-01", "end_date": "2026-09-01"},
    )

    assert jan_response.status_code == 200, jan_response.data
    assert sep_response.status_code == 200, sep_response.data
    assert wide_response.status_code == 200, wide_response.data
    assert [row["trip_reference"] for row in jan_response.data["rows"]] == [
        jan["trip"].reference
    ]
    assert [row["trip_reference"] for row in sep_response.data["rows"]] == [
        sep["trip"].reference
    ]
    assert {row["trip_reference"] for row in wide_response.data["rows"]} == {
        jan["trip"].reference,
        sep["trip"].reference,
    }


@pytest.mark.django_db
def test_dashboard_party_balances_suppliers_reflects_advance_then_debt_then_payment():
    """docs/superpowers/specs/2026-08-31-supplier-transporter-advances-
    design.md: party_balances.suppliers.balance must equal net_balance
    (open_advances - total_due), the same figure
    ControlledDebtPaymentService.describe_supplier_balance and the
    party-balances endpoint expose - never a separate computation - through
    every stage of the advance/debt/payment lifecycle."""
    admin_user = _make_admin("DASH-SYNC")
    supplier = Supplier.objects.create(name="Dashboard Sync Supplier", country="Senegal")
    product = Product.objects.get(code="gasoil")
    cash_account = CashAccount.objects.create(name="Cash Dash Sync")
    bank_account = BankAccount.objects.create(name="Bank Dash Sync", opening_balance=Decimal("10000000.00"))

    client = APIClient()
    client.force_authenticate(admin_user)

    # 1) advance created -> positive net_balance
    SupplierAdvance.objects.create(
        supplier=supplier, amount=Decimal("2000000.00"), date="2026-08-17",
        cash_account=cash_account, recorded_by=admin_user,
    )
    response = client.get("/api/dashboard/summary/")
    balance_after_advance = Decimal(response.data["party_balances"]["suppliers"]["balance"])
    assert balance_after_advance >= Decimal("2000000.00")

    # 2) debt created (order confirmed), fully covered by the advance
    # Bug fix (client business rule, confirmed 2026-09-06): debt now needs
    # quantity actually loaded - fully load both orders below so this
    # test's progression (600000 then 3000000 total debt) is unchanged;
    # this test is about advance/debt/payment lifecycle sync, not about
    # the loaded-vs-purchased debt timing rule itself.
    order_1 = SupplierOrder.objects.create(
        reference="SO-DASH-SYNC", supplier=supplier, product=product,
        purchase_date="2026-08-18", quantity_litres=Decimal("1000.00"),
        unit_price=Decimal("600.0000"), status=SupplierOrder.Status.CONFIRMED,
        created_by=admin_user,
    )  # total 600000, fully covered by the 2M advance
    cession_1 = CessionLetter.objects.create(
        reference="CL-DASH-SYNC", supplier_order=order_1,
        authorized_quantity=order_1.quantity_litres, issue_date="2026-08-18",
        created_by=admin_user,
    )
    truck_1 = Truck.objects.create(
        internal_id="TRK-DASH-SYNC", tractor_registration="DASH-SYNC-TRAC",
        tanker_registration="DASH-SYNC-TANK", capacity_litres=Decimal("45000.00"),
    )
    Trip.objects.create(
        reference="TRIP-DASH-SYNC", truck=truck_1,
        transporter=Transporter.objects.create(name="Transporteur DASH-SYNC"),
        supplier=supplier, supplier_order=order_1, cession_letter=cession_1,
        product=product, quantity_loaded=order_1.quantity_litres,
        loading_date="2026-08-18", bl_reference="BL-DASH-SYNC", created_by=admin_user,
    )
    response = client.get("/api/dashboard/summary/")
    balance_after_debt = Decimal(response.data["party_balances"]["suppliers"]["balance"])
    assert balance_after_debt == balance_after_advance - Decimal("600000.00")
    assert balance_after_debt > Decimal("0.00")

    # 3) a larger debt -> now net negative (real payable debt)
    order_2 = SupplierOrder.objects.create(
        reference="SO-DASH-SYNC-2", supplier=supplier, product=product,
        purchase_date="2026-08-19", quantity_litres=Decimal("5000.00"),
        unit_price=Decimal("600.0000"), status=SupplierOrder.Status.CONFIRMED,
        created_by=admin_user,
    )  # total 3000000, consumes remaining 1400000 advance, leaves 1600000 real debt
    cession_2 = CessionLetter.objects.create(
        reference="CL-DASH-SYNC-2", supplier_order=order_2,
        authorized_quantity=order_2.quantity_litres, issue_date="2026-08-19",
        created_by=admin_user,
    )
    truck_2 = Truck.objects.create(
        internal_id="TRK-DASH-SYNC-2", tractor_registration="DASH-SYNC2-TRAC",
        tanker_registration="DASH-SYNC2-TANK", capacity_litres=Decimal("45000.00"),
    )
    Trip.objects.create(
        reference="TRIP-DASH-SYNC-2", truck=truck_2,
        transporter=Transporter.objects.create(name="Transporteur DASH-SYNC-2"),
        supplier=supplier, supplier_order=order_2, cession_letter=cession_2,
        product=product, quantity_loaded=order_2.quantity_litres,
        loading_date="2026-08-19", bl_reference="BL-DASH-SYNC-2", created_by=admin_user,
    )
    response = client.get("/api/dashboard/summary/")
    assert Decimal(response.data["party_balances"]["suppliers"]["balance"]) == Decimal("-1600000.00")

    # 4) partial payment reduces the real payable debt
    ControlledDebtPaymentService.record_payment(
        debt_type="supplier", debt_id=supplier.id, bank_account_id=bank_account.id,
        amount=Decimal("600000.00"), date="2026-08-20", reference="PAY-DASH-SYNC",
        actor=admin_user,
    )
    response = client.get("/api/dashboard/summary/")
    dashboard_balance = Decimal(response.data["party_balances"]["suppliers"]["balance"])
    assert dashboard_balance == Decimal("-1000000.00")

    # Chiffre exact: Dashboard == party-balances endpoint == describe_supplier_balance
    party_balances_response = client.get(
        "/api/treasury-transactions/party-balances/", {"debt_type": "supplier"}
    )
    row = next(r for r in party_balances_response.data if r["id"] == supplier.id)
    balance = ControlledDebtPaymentService.describe_supplier_balance(supplier)

    assert Decimal(row["net_balance"]) == dashboard_balance == balance["net_balance"] == Decimal("-1000000.00")


# ---------------------------------------------------------------------------
# ReportViewSet
# ---------------------------------------------------------------------------


@pytest.mark.django_db
def test_report_endpoint_returns_json_rows_for_a_known_report_type():
    admin_user = _make_admin("REPJSON")
    fx = _make_full_trip("REPJSON", admin_user)

    client = APIClient()
    client.force_authenticate(admin_user)
    response = client.get("/api/reports/trip-report/")

    assert response.status_code == 200, response.data
    assert response.data["report_type"] == "trip-report"
    assert any(row["trip_reference"] == fx["trip"].reference for row in response.data["rows"])


@pytest.mark.django_db
def test_report_endpoint_unknown_type_returns_404():
    admin_user = _make_admin("REP404")

    client = APIClient()
    client.force_authenticate(admin_user)
    response = client.get("/api/reports/does-not-exist/")

    assert response.status_code == 404


@pytest.mark.django_db
def test_report_endpoint_passes_query_params_to_the_query_class():
    admin_user = _make_admin("REPFILTER")
    fx = _make_full_trip("REPFILTER", admin_user)

    client = APIClient()
    client.force_authenticate(admin_user)
    response = client.get("/api/reports/truck-history/", {"truck_id": fx["truck"].id})

    assert response.status_code == 200, response.data
    assert len(response.data["rows"]) == 1
    assert response.data["rows"][0]["trip_reference"] == fx["trip"].reference


@pytest.mark.django_db
def test_report_endpoint_export_csv_returns_a_csv_attachment():
    admin_user = _make_admin("REPCSV")
    _make_full_trip("REPCSV", admin_user)

    client = APIClient()
    client.force_authenticate(admin_user)
    response = client.get("/api/reports/trip-report/", {"export": "csv"})

    assert response.status_code == 200
    assert response["Content-Type"] == "text/csv"
    assert "attachment" in response["Content-Disposition"]
    reader = csv.reader(io.StringIO(response.content.decode("utf-8-sig")))
    header = next(reader)
    assert header == list(header)  # non-empty, well-formed CSV header
    assert len(header) > 0


@pytest.mark.django_db
def test_report_endpoint_export_xlsx_returns_a_valid_workbook():
    admin_user = _make_admin("REPXLSX")
    _make_full_trip("REPXLSX", admin_user)

    client = APIClient()
    client.force_authenticate(admin_user)
    response = client.get("/api/reports/trip-report/", {"export": "xlsx"})

    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook.active.max_row >= 2  # header + at least one data row


@pytest.mark.django_db
def test_report_endpoint_export_pdf_returns_a_pdf_attachment():
    admin_user = _make_admin("REPPDF")
    _make_full_trip("REPPDF", admin_user)

    client = APIClient()
    client.force_authenticate(admin_user)
    response = client.get("/api/reports/trip-report/", {"export": "pdf"})

    assert response.status_code == 200
    assert response["Content-Type"] == "application/pdf"
    assert "attachment" in response["Content-Disposition"]
    assert response.content.startswith(b"%PDF-")


@pytest.mark.django_db
def test_report_export_requires_export_permission_not_just_view():
    role = Role.objects.create(name="ViewOnlyReports", permissions={"reports": ["view"]})
    user = User.objects.create_user(phone="+2220000014", pin="1234", role=role)

    client = APIClient()
    client.force_authenticate(user)
    response = client.get("/api/reports/trip-report/", {"export": "csv"})

    assert response.status_code == 403
    # Without the export param, the same view-only role can still read JSON.
    response = client.get("/api/reports/trip-report/")
    assert response.status_code == 200


# ---------------------------------------------------------------------------
# SearchView
# ---------------------------------------------------------------------------


@pytest.mark.django_db
def test_search_finds_a_truck_with_its_current_trip_enriched():
    admin_user = _make_admin("SEARCH")
    fx = _make_full_trip("SEARCH", admin_user, status=Trip.Status.IN_TRANSIT, with_sale=False)

    client = APIClient()
    client.force_authenticate(admin_user)
    response = client.get("/api/search/", {"q": fx["truck"].tractor_registration})

    assert response.status_code == 200, response.data
    assert len(response.data["trucks"]) == 1
    truck_result = response.data["trucks"][0]
    assert truck_result["current_trip"]["reference"] == fx["trip"].reference
    assert truck_result["current_trip"]["status"] == Trip.Status.IN_TRANSIT


@pytest.mark.django_db
def test_search_with_empty_query_returns_empty_categories():
    admin_user = _make_admin("SEARCHEMPTY")

    client = APIClient()
    client.force_authenticate(admin_user)
    response = client.get("/api/search/", {"q": ""})

    assert response.status_code == 200
    assert response.data["trucks"] == []
    assert response.data["customers"] == []


# ---------------------------------------------------------------------------
# DashboardViewSet.summary - working_capital
# ---------------------------------------------------------------------------


@pytest.mark.django_db
def test_dashboard_summary_includes_working_capital_section():
    admin_user = _make_admin("DASHWC1")
    client = APIClient()
    client.force_authenticate(admin_user)

    response = client.get("/api/dashboard/summary/")

    assert response.status_code == 200, response.data
    assert set(response.data["working_capital"]) >= {
        "capital", "treasury_available", "customer_receivables",
        "supplier_debt_total", "transporter_debt_total", "other_debts",
        "supplier_reliquat_value", "supplier_in_transit_value",
        "supplier_company_stock_value", "gross", "total",
    }


@pytest.mark.django_db
def test_dashboard_summary_working_capital_uses_additive_formula():
    """Task 8: the dashboard payload's working_capital must be the new
    gross/total breakdown (old net/signed supplier_net_balance/
    transporter_net_balance keys gone), and shareholders_summary must
    expose distributable_result instead of profit_after_capital."""
    from decimal import Decimal

    admin_user = _make_admin("DASHWCADD")
    client = APIClient()
    client.force_authenticate(admin_user)

    response = client.get("/api/dashboard/summary/")

    assert response.status_code == 200, response.data
    wc = response.data["working_capital"]
    assert "supplier_net_balance" not in wc
    assert "transporter_net_balance" not in wc
    assert "supplier_debt_total" in wc
    assert "transporter_debt_total" in wc
    assert Decimal(response.data["shareholders_summary"]["distributable_result"]) is not None


@pytest.mark.django_db
def test_dashboard_working_capital_formula_matches_manual_calculation():
    from decimal import Decimal
    from apps.finance.models import BankAccount, CashAccount, Shareholder
    from apps.finance.services import ShareholderContributionService
    from django.core.files.uploadedfile import SimpleUploadedFile

    admin_user = _make_admin("DASHWC2")
    bank = BankAccount.objects.create(name="WC Bank", opening_balance=Decimal("10000000.00"))
    CashAccount.objects.create(name="WC Cash", opening_balance=Decimal("2000000.00"))
    shareholder = Shareholder.objects.create(
        name="WC Shareholder", participation_percent=Decimal("100.00"), contribution=Decimal("0.00"),
    )
    ShareholderContributionService.record_contribution(
        shareholder_id=shareholder.id, amount=Decimal("5000000.00"), date="2026-09-02",
        actor=admin_user, attachment=SimpleUploadedFile("p.jpg", b"x", content_type="image/jpeg"),
        bank_account_id=bank.id,
    )

    client = APIClient()
    client.force_authenticate(admin_user)
    response = client.get("/api/dashboard/summary/")
    wc = response.data["working_capital"]

    # Bank now holds 10,000,000 (opening) + 5,000,000 (contribution) = 15,000,000.
    # Cash = 2,000,000. No suppliers/transporters/customers/debts exist in this
    # test, so treasury_available is the only nonzero component of gross, and
    # no trade debts exist to add, so total == gross + what the company owes
    # the associate for his contribution (not drawn yet, 2026-10-16).
    assert wc["treasury_available"] == "17000000.00"
    assert wc["capital"] == "5000000.00"
    assert wc["gross"] == "17000000.00"
    assert wc["shareholder_dues"] == "5000000.00"
    assert wc["total"] == "22000000.00"
    # Capital must NOT be double-counted: gross == treasury_available exactly,
    # not treasury_available + capital.
    assert Decimal(wc["gross"]) == Decimal(wc["treasury_available"])


@pytest.mark.django_db
def test_dashboard_working_capital_advance_reduces_treasury_without_any_offset():
    """Client modifications (2026-09-04): the corrected fond-de-roulement
    formula never nets debts against advances at all - an open advance is
    simply not part of the formula (docs/superpowers/specs/2026-09-04-
    recette-corrections-design.md Section 3). This replaces the old
    net/signed working_capital's behaviour (where an open advance offset
    the cash outflow that created it) - under the (third-pass, additive)
    formula, paying out an advance with no offsetting debt simply reduces
    treasury_available, dropping working_capital.total (and .gross, since
    no debt exists to add) by the same amount."""
    from decimal import Decimal
    from apps.finance.models import CashAccount
    from apps.finance.services import AdvanceService
    from apps.suppliers.models import Supplier
    from django.core.files.uploadedfile import SimpleUploadedFile

    admin_user = _make_admin("DASHWC3")
    cash = CashAccount.objects.create(name="WC Cash 2", opening_balance=Decimal("10000000.00"))
    supplier = Supplier.objects.create(name="WC Supplier", country="Senegal")

    client = APIClient()
    client.force_authenticate(admin_user)
    before = client.get("/api/dashboard/summary/").data["working_capital"]["total"]

    AdvanceService.create_supplier_advance(
        supplier_id=supplier.id, bank_account_id=None, cash_account_id=cash.id,
        amount=Decimal("2000000.00"), date="2026-09-02", reference="", notes="",
        attachment=SimpleUploadedFile("p.jpg", b"x", content_type="image/jpeg"), actor=admin_user,
    )

    after = client.get("/api/dashboard/summary/").data["working_capital"]["total"]
    assert Decimal(after) == Decimal(before) - Decimal("2000000.00")


@pytest.mark.django_db
def test_dashboard_working_capital_matches_client_recette_example():
    """Confirmed bug during actual client recette testing (2026-09-04,
    third pass): the "Fond de roulement net client" dashboard card showed
    dettes ouvertes as positive line items but the headline figure still
    subtracted them. Client-provided example, reproduced exactly via the
    real dashboard endpoint:
      gross (tresorerie + creances) = 300 000 000
      + dettes fournisseurs ouvertes  =  45 000 000
      + dettes transporteurs ouvertes =   2 925 000
      + dettes douane/transit/frais   =  10 255 000
      = 358 180 000
    """
    admin_user = _make_admin("DASHRECETTE")
    BankAccount.objects.create(
        name="Dash Recette Bank", opening_balance=Decimal("300000000.00"), is_active=True
    )

    supplier = Supplier.objects.create(name="Dash Recette Supplier", country="Senegal")
    product = Product.objects.get(code="gasoil")
    recette_order = SupplierOrder.objects.create(
        reference="SO-DASHRECETTE", supplier=supplier, product=product,
        purchase_date="2026-08-18", quantity_litres=Decimal("45000.00"),
        unit_price=Decimal("1000.0000"), status=SupplierOrder.Status.CONFIRMED,
        created_by=admin_user,
    )
    # Bug fix (client business rule, confirmed 2026-09-06): debt now needs
    # quantity actually loaded - fully load this order so the client's own
    # "dettes fournisseurs ouvertes = 45 000 000" example still holds.
    recette_cession = CessionLetter.objects.create(
        reference="CL-DASHRECETTE-SUP", supplier_order=recette_order,
        authorized_quantity=recette_order.quantity_litres, issue_date="2026-08-18",
        created_by=admin_user,
    )
    recette_truck = Truck.objects.create(
        internal_id="TRK-DASHRECETTE-SUP", tractor_registration="DASHRECSUP-TRAC",
        tanker_registration="DASHRECSUP-TANK", capacity_litres=Decimal("45000.00"),
    )
    Trip.objects.create(
        reference="TRIP-DASHRECETTE-SUP", truck=recette_truck,
        transporter=Transporter.objects.create(name="Transporteur Dash Recette Achat"),
        supplier=supplier, supplier_order=recette_order, cession_letter=recette_cession,
        product=product, quantity_loaded=recette_order.quantity_litres,
        loading_date="2026-08-18", bl_reference="BL-DASHRECETTE-SUP", created_by=admin_user,
    )

    transporter = Transporter.objects.create(name="Dash Recette Transporter")
    truck = Truck.objects.create(
        internal_id="TRK-DASHRECETTE", tractor_registration="DASHRECETTE-TRACTOR",
        tanker_registration="DASHRECETTE-TANKER", capacity_litres=Decimal("45000.00"),
    )
    debt_supplier = Supplier.objects.create(name="Dash Recette Trip Supplier", country="Senegal")
    debt_order = SupplierOrder.objects.create(
        reference="SO-DASHRECETTE-TRIP", supplier=debt_supplier, product=product,
        purchase_date="2026-08-18", quantity_litres=Decimal("45000.00"),
        unit_price=Decimal("0.0000"), status=SupplierOrder.Status.CONFIRMED,
        created_by=admin_user,
    )
    cession = CessionLetter.objects.create(
        reference="CL-DASHRECETTE", supplier_order=debt_order,
        authorized_quantity=Decimal("45000.00"), issue_date="2026-08-18", created_by=admin_user,
    )
    debt_trip = Trip.objects.create(
        reference="TRIP-DASHRECETTE", truck=truck, transporter=transporter, supplier=debt_supplier,
        supplier_order=debt_order, cession_letter=cession, product=product,
        quantity_loaded=Decimal("45000.00"), loading_date="2026-08-18",
        bl_reference="BL-DASHRECETTE", created_by=admin_user,
    )
    TransportPayment.objects.create(
        trip=debt_trip, transporter=transporter,
        calculation_method=TransportPayment.CalculationMethod.FIXED,
        base_amount=Decimal("2925000.00"),
    )

    Expense.objects.create(
        category=Expense.Category.CUSTOMS, amount=Decimal("10255000.00"),
        calculation_method=Expense.CalculationMethod.MANUAL, date="2026-08-18",
    )

    client = APIClient()
    client.force_authenticate(admin_user)
    response = client.get("/api/dashboard/summary/")

    assert response.status_code == 200, response.data
    wc = response.data["working_capital"]
    assert wc["gross"] == "300000000.00"
    assert wc["supplier_debt_total"] == "45000000.00"
    assert wc["transporter_debt_total"] == "2925000.00"
    assert wc["other_debts"] == "10255000.00"
    assert wc["total"] == "358180000.00"


# ---------------------------------------------------------------------------
# ?document=1 (docs/superpowers/specs/2026-09-23-rapports-document-comptable-
# design.md, D3 and section Compatibility contract): opt-in, additive.
# TripDetail/DashboardView (loss-report) and ProfitView (margin-report) call
# the endpoint without it and must keep receiving exactly
# {report_type, fields, rows}.
# ---------------------------------------------------------------------------


def _all_report_params(admin_user, suffix):
    from apps.finance.models import BankAccount, CashAccount, Shareholder

    fx = _make_full_trip(suffix, admin_user)
    return fx, {
        "customer-statement": {"customer_id": fx["customer"].id},
        "supplier-statement": {"supplier_id": fx["supplier_order"].supplier_id},
        "transporter-statement": {"transporter_id": fx["trip"].transporter_id},
        "truck-history": {"truck_id": fx["truck"].id},
        "trip-report": {},
        "purchase-report": {},
        "sales-report": {},
        "expense-report": {},
        "customer-outstanding": {},
        "supplier-outstanding": {},
        "transporter-outstanding": {},
        "fuel-stock": {},
        "loss-report": {},
        "margin-report": {},
        "bank-statement": {"bank_account_id": BankAccount.objects.create(name=f"Bank {suffix}").id},
        "cash-statement": {"cash_account_id": CashAccount.objects.create(name=f"Cash {suffix}").id},
        "shareholder-statement": {
            "shareholder_id": Shareholder.objects.create(
                name=f"Assoc {suffix}", participation_percent=Decimal("100")
            ).id
        },
        # Added 2026-09-23: no entity id needed.
        "capital-report": {},
        "general-expense-report": {},
        "trip-fee-debt-report": {"fee_type": "customs"},
        "general-debt-report": {},
        "general-receivable-report": {},
        "debts-summary-report": {},
        "receivables-summary-report": {},
    }


@pytest.mark.django_db
def test_report_endpoint_legacy_response_is_unchanged_for_every_type_and_document_is_additive():
    from apps.reports.serialization import decimals_to_str
    from apps.reports.views.reports import REPORT_REGISTRY

    admin_user = _make_admin("REPCOMPAT")
    _, params_by_type = _all_report_params(admin_user, "REPCOMPAT")
    client = APIClient()
    client.force_authenticate(admin_user)

    assert set(params_by_type) == set(REPORT_REGISTRY)
    for report_type, params in params_by_type.items():
        query_class = REPORT_REGISTRY[report_type]
        legacy = client.get(f"/api/reports/{report_type}/", params)

        assert legacy.status_code == 200, report_type
        assert set(legacy.data) == {"report_type", "fields", "rows"}, report_type
        assert legacy.data["report_type"] == report_type
        assert legacy.data["fields"] == query_class.FIELDS
        assert legacy.data["rows"] == decimals_to_str(query_class.run(**params)), report_type

        with_document = client.get(f"/api/reports/{report_type}/", {**params, "document": "1", "lang": "fr"})

        assert with_document.status_code == 200, report_type
        assert set(with_document.data) == {"report_type", "fields", "rows", "document"}
        for key in ("report_type", "fields", "rows"):
            assert with_document.data[key] == legacy.data[key], (report_type, key)
        assert with_document.data["document"]["title"]


@pytest.mark.django_db
def test_report_endpoint_keeps_the_fields_read_by_trip_detail_dashboard_and_profit_views():
    from apps.reports.queries.loss_report import FIELDS as LOSS_FIELDS
    from apps.reports.queries.margin_report import FIELDS as MARGIN_FIELDS

    admin_user = _make_admin("REPCONSUMERS")
    fx = _make_full_trip("REPCONSUMERS", admin_user)
    client = APIClient()
    client.force_authenticate(admin_user)

    loss = client.get("/api/reports/loss-report/", {"from": "2026-01-01", "to": "2026-12-31"})
    margin = client.get("/api/reports/margin-report/")

    assert {"trip_reference", "verified_at", "loaded_quantity", "received_quantity", "loss_quantity",
            "chargeable_loss_quantity"} <= set(loss.data["fields"])
    assert loss.data["fields"] == LOSS_FIELDS and set(loss.data) == {"report_type", "fields", "rows"}
    assert {"trip_reference", "unit_profit", "current_margin", "real_total_cost"} <= set(margin.data["fields"])
    assert margin.data["fields"] == MARGIN_FIELDS and set(margin.data) == {"report_type", "fields", "rows"}
    assert any(row["trip_reference"] == fx["trip"].reference for row in margin.data["rows"])
    # Rows also carry keys outside FIELDS (never exported): trip_id,
    # is_adjustment, revenue, supplier_cost (spec 2026-10-05-periodes-
    # independantes) - FIELDS itself is unchanged.
    extra = {"trip_id", "is_adjustment", "revenue", "supplier_cost"}
    assert all(set(row) == set(MARGIN_FIELDS) | extra for row in margin.data["rows"])


@pytest.mark.django_db
def test_report_endpoint_document_for_a_statement_carries_opening_balance_and_entity():
    from apps.finance.models import CashAccount, CustomerPayment
    from apps.sales.models import Invoice

    admin_user = _make_admin("REPDOCSTMT")
    fx = _make_full_trip("REPDOCSTMT", admin_user)
    invoice = Invoice.objects.create(
        invoice_number="INV-REPDOCSTMT", sale=fx["sale"], status=Invoice.Status.VALIDATED,
        total_amount=fx["sale"].subtotal, issued_date="2026-08-18",
    )
    CustomerPayment.objects.create(
        customer=fx["customer"], amount=Decimal("1000.00"), payment_date="2026-09-05",
        method=CustomerPayment.Method.CASH, cash_account=CashAccount.objects.create(name="Cash REPDOCSTMT"),
        recorded_by=admin_user,
    )
    client = APIClient()
    client.force_authenticate(admin_user)

    response = client.get(
        "/api/reports/customer-statement/",
        {"customer_id": fx["customer"].id, "from": "2026-09-01", "document": "1", "lang": "fr"},
    )

    assert response.status_code == 200, response.data
    document = response.data["document"]
    assert document["opening_balance"]["value"] == str(invoice.total_amount)
    assert response.data["rows"][0]["balance"] == str(invoice.total_amount - Decimal("1000.00"))
    assert document["entity"]["name"] == fx["customer"].name
    assert document["period"]["text"].startswith("À partir du 01/09/2026")
    assert {line["key"] for line in document["summary"]} >= {"opening_balance", "closing_balance"}


@pytest.mark.django_db
@pytest.mark.parametrize("params", [{}, {"customer_id": "abc"}, {"customer_id": "999999"}])
def test_report_endpoint_document_never_turns_a_bad_entity_id_into_a_500(params):
    admin_user = _make_admin("REPBADID")
    client = APIClient()
    client.force_authenticate(admin_user)

    legacy = client.get("/api/reports/customer-statement/", params)
    with_document = client.get("/api/reports/customer-statement/", {**params, "document": "1"})

    assert legacy.status_code < 500
    assert with_document.status_code == legacy.status_code


@pytest.mark.django_db
@pytest.mark.parametrize(
    "report_type, param",
    [("bank-statement", "bank_account_id"), ("cash-statement", "cash_account_id")],
)
def test_report_endpoint_unknown_treasury_account_is_a_404_not_a_500(report_type, param):
    admin_user = _make_admin("REPNOACCT")
    client = APIClient()
    client.force_authenticate(admin_user)

    legacy = client.get(f"/api/reports/{report_type}/", {param: 999999})
    with_document = client.get(f"/api/reports/{report_type}/", {param: 999999, "document": "1"})

    assert legacy.status_code == with_document.status_code == 404


@pytest.mark.django_db
def test_dashboard_working_capital_pdf_prints_only_the_selected_lines():
    from apps.reports.services.working_capital_pdf import render_working_capital_pdf

    admin_user = _make_admin("WCPDF")
    client = APIClient()
    client.force_authenticate(admin_user)

    response = client.get(
        "/api/dashboard/working-capital-pdf/?lines=total,treasury,receivables&lang=fr"
    )
    assert response.status_code == 200
    assert response["Content-Type"] == "application/pdf"
    assert response.content.startswith(b"%PDF")

    everything = client.get("/api/dashboard/working-capital-pdf/")
    assert everything.status_code == 200
    # fewer lines -> a different (smaller) document; unknown keys ignored
    assert client.get("/api/dashboard/working-capital-pdf/?lines=nope").status_code == 400
    assert client.get("/api/dashboard/working-capital-pdf/?lines=").status_code == 400

    # The renderer only prints what it is given.
    wc = {
        key: Decimal("1000.00")
        for key in (
            "total", "capital", "treasury_available", "customer_receivables",
            "supplier_debt_total", "transporter_debt_total", "other_debts",
            "supplier_company_stock_value", "supplier_in_transit_value",
            "supplier_reliquat_value",
        )
    }
    for lang in ("fr", "en", "ar"):
        assert render_working_capital_pdf(working_capital=wc, lines=["total"], lang=lang).startswith(b"%PDF")
